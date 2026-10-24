from openpyxl import Workbook
import pytest

from antecipe.adapters.readers import COLUMNS, read_rows, parse_row, FileError


def values():
    return [
        "Fundo",
        "92.234.400/0001-80",
        "Cedente",
        "10.062.050/0001-57",
        "Sacado",
        "22.144.961/0001-02",
        "000351153",
        "Duplicata de Serviço",
        "10000,00",
        "31/10/2026",
        "01/10/2026",
    ]


def csv_bytes(rows=None):
    return (";".join(COLUMNS) + "\n" + "\n".join(";".join(r) for r in (rows or [values()])) + "\n").encode(
        "utf-8-sig"
    )


def test_csv_xlsx_equivalent_physical_rows(tmp_path):
    csv = tmp_path / "a.csv"
    csv.write_bytes(csv_bytes([values(), [""] * 11, values()]))
    xlsx = tmp_path / "a.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(COLUMNS)
    ws.append(values())
    ws.append([""] * 11)
    ws.append(values())
    wb.save(xlsx)
    a = list(read_rows(csv, "CSV"))
    b = list(read_rows(xlsx, "XLSX"))
    assert a == b
    assert [line for line, raw in a] == [2, 4]
    parsed = parse_row(a[0][1])
    assert parsed["nominal"] == "10000.00" and parsed["errors"] == []
    assert parsed["raw"]["NU_DOCUMENTO"] == "000351153"


def test_data_error_pl_preservation_and_unreadable_nominal():
    raw = dict(zip(COLUMNS, values()))
    raw["DATA_AQUISICAO"] = "31/02/2026"
    parsed = parse_row(raw)
    assert parsed["nominal"] == "10000.00" and not parsed["blocks_pl"]
    assert parsed["errors"][0]["codigo"] == "DATA_INVALIDA"
    raw["VALOR_NOMINAL"] = "NaN"
    assert parse_row(raw)["blocks_pl"]
    raw["DOC_FUNDO"] = "ABC"
    assert parse_row(raw)["fund_document"] is None


def test_structural_failure_and_sheet_formula(tmp_path):
    path = tmp_path / "bad.csv"
    path.write_text("one;two\nx;y", encoding="utf8")
    with pytest.raises(FileError):
        list(read_rows(path, "CSV"))
    path = tmp_path / "a.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Notas"
    ws.append(COLUMNS)
    ws.append(values())
    ws["I2"] = "=100+1"
    wb.save(path)
    raw = list(read_rows(path, "XLSX", "Notas"))[0][1]
    assert parse_row(raw)["blocks_pl"]
    assert parse_row(raw)["errors"][0]["codigo"] == "FORMULA_SEM_VALOR"
    with pytest.raises(FileError):
        list(read_rows(path, "XLSX", "Ausente"))
    path.write_bytes(b"corrupted")
    with pytest.raises(FileError):
        list(read_rows(path, "XLSX"))


def test_xlsx_fallback_reads_uuid_storage_path(tmp_path):
    path = tmp_path / "550e8400-e29b-41d4-a716-446655440000"
    wb = Workbook()
    ws = wb.active
    ws.append(COLUMNS)
    ws.append(values())
    ws["I2"] = "=100+1"
    wb.save(path)
    raw = list(read_rows(path, "XLSX"))[0][1]
    parsed = parse_row(raw)
    assert parsed["errors"][0]["codigo"] == "FORMULA_SEM_VALOR"
    assert parsed["raw"]["VALOR_NOMINAL"] == "=100+1"


def test_zip_without_workbook_is_structural_failure(tmp_path):
    from zipfile import ZipFile

    path = tmp_path / "bad.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr("xl/not-a-workbook.xml", "bad")
    with pytest.raises(FileError):
        list(read_rows(path, "XLSX"))


def test_csv_quoted_newline_preserves_starting_physical_line(tmp_path):
    import csv

    path = tmp_path / "multiline.csv"
    first = values()
    first[2] = "Cedente\nFilial"
    with path.open("w", newline="", encoding="utf8") as stream:
        writer = csv.writer(stream, delimiter=";")
        writer.writerows([COLUMNS, first, values()])
    rows = list(read_rows(path, "CSV"))
    assert [line for line, _ in rows] == [2, 4]
    assert rows[0][1]["NOME_CEDENTE"] == "Cedente\nFilial"
