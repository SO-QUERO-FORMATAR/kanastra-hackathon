import csv
import re
from contextlib import contextmanager
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from python_calamine import CalamineWorkbook, CalamineError

from antecipe.domain.entities import normalize_document, document_root, Title, Fund, Origin
from antecipe.domain.pricing import money

COLUMNS = [
    "NOME_FUNDO",
    "DOC_FUNDO",
    "NOME_CEDENTE",
    "DOC_CEDENTE",
    "NOME_SACADO",
    "DOC_SACADO",
    "NU_DOCUMENTO",
    "TIPO_RECEBIVEL",
    "VALOR_NOMINAL",
    "DATA_VENCIMENTO_ORIGINAL",
    "DATA_AQUISICAO",
]
FORMULA = "__ANTECIPE_FORMULA_SEM_VALOR__"
NUMERIC_ID = "__ANTECIPE_IDENTIFICADOR_NUMERICO__"


class FileError(ValueError):
    pass


@contextmanager
def open_xlsx(path, data_only):
    with path.open("rb") as stream:
        book = load_workbook(stream, read_only=True, data_only=data_only)
        try:
            yield book
        finally:
            book.close()


def _text(value, column):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, (int, float)):
        if column.startswith("DOC_") or column == "NU_DOCUMENTO":
            return NUMERIC_ID + str(value)
        return str(value).replace(".", ",") if column == "VALOR_NOMINAL" else str(value)
    return str(value)


def read_calamine(path, sheet):
    with CalamineWorkbook.from_path(path) as workbook:
        if sheet and sheet not in workbook.sheet_names:
            raise FileError("ESTRUTURA_ARQUIVO_INVALIDA: aba ausente")
        names = [sheet] if sheet else workbook.sheet_names
        selected = next(
            (
                workbook.get_sheet_by_name(name)
                for name in names
                if workbook.get_sheet_by_name(name).start is not None
            ),
            None,
        )
        if selected is None:
            raise FileError("ESTRUTURA_ARQUIVO_INVALIDA: aba vazia")
        if selected.start != (0, 0):
            raise FileError("ESTRUTURA_ARQUIVO_INVALIDA: cabeçalho deve começar em A1")
        rows = selected.iter_rows()
        header = next(rows, [])
        if len(header) != len(set(header)) or not set(COLUMNS) <= set(header):
            raise FileError("ESTRUTURA_ARQUIVO_INVALIDA: onze colunas obrigatórias")
        for line, cells in enumerate(rows, 2):
            if not any(v not in ("", None) for v in cells):
                continue
            row = {col: _text(value, col) for col, value in zip(header, cells) if col in COLUMNS}
            yield line, {col: row.get(col) for col in COLUMNS}


def read_rows(path: Path, format: str, sheet: str | None = None):
    try:
        if format == "CSV":
            with path.open(encoding="utf-8-sig", newline="") as stream:
                reader = csv.reader(stream, delimiter=";", strict=True)
                header = next(reader, [])
                if len(header) != len(set(header)) or not set(COLUMNS) <= set(header):
                    raise FileError("ESTRUTURA_ARQUIVO_INVALIDA: onze colunas obrigatórias")
                previous_line = reader.line_num
                for cells in reader:
                    line = previous_line + 1
                    previous_line = reader.line_num
                    if not any(v.strip() for v in cells):
                        continue
                    if len(cells) != len(header):
                        raise FileError("ESTRUTURA_ARQUIVO_INVALIDA: largura da linha")
                    row = dict(zip(header, cells))
                    yield line, {k: row[k] for k in COLUMNS}
        elif format == "XLSX":
            # Read-only fallback keeps 500k sheets bounded and exposes uncached formulas safely.
            with ZipFile(path) as archive:
                if any("vbaProject" in n for n in archive.namelist()):
                    raise FileError("ARQUIVO_CORROMPIDO: macros não suportadas")
                # Calamine materializes the sheet internally. Bound uncompressed input before selecting it.
                small = sum(info.file_size for info in archive.infolist()) <= 16 * 1024 * 1024
                formulas = small and any(
                    re.search(rb"<(?:[a-zA-Z]+:)?f(?:\s|>)", archive.read(n))
                    for n in archive.namelist()
                    if n.startswith("xl/worksheets/") and n.endswith(".xml")
                )
            if small and not formulas:
                yield from read_calamine(path, sheet)
                return
            with open_xlsx(path, False) as wb, open_xlsx(path, True) as cached:
                if sheet and sheet not in wb.sheetnames:
                    raise FileError("ESTRUTURA_ARQUIVO_INVALIDA: aba ausente")
                ws = (
                    wb[sheet]
                    if sheet
                    else next(
                        (
                            s
                            for s in wb
                            if any(
                                v is not None for row in s.iter_rows(max_row=1, values_only=True) for v in row
                            )
                        ),
                        wb.worksheets[0],
                    )
                )
                rows, cached_rows = ws.iter_rows(), cached[ws.title].iter_rows()
                header = [c.value for c in next(rows, [])]
                next(cached_rows, None)
                if len(header) != len(set(header)) or not set(COLUMNS) <= set(header):
                    raise FileError("ESTRUTURA_ARQUIVO_INVALIDA: onze colunas obrigatórias")
                for line, (cells, values) in enumerate(zip(rows, cached_rows), 2):
                    if not any(c.value is not None for c in cells):
                        continue
                    raw = {}
                    for col, cell, cache in zip(header, cells, values):
                        if col not in COLUMNS:
                            continue
                        value = cell.value
                        if cell.data_type == "f":
                            value = cache.value if cache.value is not None else FORMULA + str(cell.value)
                        raw[col] = _text(value, col)
                    yield line, {k: raw.get(k) for k in COLUMNS}
        else:
            raise FileError("TIPO_ARQUIVO_NAO_SUPORTADO")
    except FileError:
        raise
    except (UnicodeError, csv.Error, BadZipFile, OSError, ValueError, KeyError, CalamineError) as exc:
        raise FileError("ARQUIVO_CORROMPIDO: " + str(exc)) from exc


def parse_row(raw: dict) -> dict:
    errors: list[dict[str, str]] = []
    parsed: dict = {"raw": raw, "errors": errors, "nominal": None, "fund_document": None, "blocks_pl": False}
    for key in COLUMNS:
        value = raw.get(key)
        code = None
        if isinstance(value, str) and value.startswith(FORMULA):
            code = "FORMULA_SEM_VALOR"
        elif value is None or value == "":
            code = "OBRIGATORIO"
        elif value.startswith(NUMERIC_ID):
            code = "FORMATO_INVALIDO"
        if code:
            errors.append({"campo": key, "codigo": code, "mensagem": f"{key}: {code}"})
    bad = {e["campo"] for e in errors}
    if "DOC_FUNDO" not in bad:
        doc = normalize_document(raw["DOC_FUNDO"])
        if document_root(doc) is None:
            errors.append(
                {
                    "campo": "DOC_FUNDO",
                    "codigo": "FORMATO_INVALIDO",
                    "mensagem": "Fundo deve conter 14 dígitos",
                }
            )
        else:
            parsed["fund_document"] = doc
    if "VALOR_NOMINAL" not in bad:
        value = raw["VALOR_NOMINAL"].strip()
        try:
            if not re.fullmatch(r"-?(?:[0-9]+|[0-9]{1,3}(?:\.[0-9]{3})+)(?:,[0-9]{1,2})?", value):
                raise ValueError("Decimal inválido")
            number = Decimal(value.replace(".", "").replace(",", "."))
            if len(number.as_tuple().digits) > 36:
                raise ValueError("Precisão excede NUMERIC(38,2)")
            parsed["nominal"] = money(number)
        except (ValueError, InvalidOperation):
            errors.append(
                {"campo": "VALOR_NOMINAL", "codigo": "DECIMAL_INVALIDO", "mensagem": "Valor decimal inválido"}
            )
    for key in ["DATA_AQUISICAO", "DATA_VENCIMENTO_ORIGINAL"]:
        if key not in bad:
            try:
                parsed[key] = datetime.strptime(raw[key], "%d/%m/%Y").date().isoformat()
            except ValueError:
                errors.append(
                    {
                        "campo": key,
                        "codigo": "DATA_INVALIDA",
                        "mensagem": "Data inexistente ou formato inválido",
                    }
                )
    parsed["blocks_pl"] = parsed["nominal"] is None or parsed["fund_document"] is None
    parsed["raw"] = {
        key: value.removeprefix(NUMERIC_ID).removeprefix(FORMULA) if isinstance(value, str) else value
        for key, value in raw.items()
    }
    return parsed


def as_title(parsed: dict, fund_id: str, origin: Origin) -> Title:
    raw = parsed["raw"]
    return Title(
        Fund(fund_id, raw["NOME_FUNDO"], parsed["fund_document"]),
        origin,
        raw["NOME_CEDENTE"],
        raw["DOC_CEDENTE"],
        raw["NOME_SACADO"],
        raw["DOC_SACADO"],
        raw["NU_DOCUMENTO"],
        raw["TIPO_RECEBIVEL"],
        Decimal(parsed["nominal"]),
        date.fromisoformat(parsed["DATA_AQUISICAO"]),
        date.fromisoformat(parsed["DATA_VENCIMENTO_ORIGINAL"]),
        raw,
    )
